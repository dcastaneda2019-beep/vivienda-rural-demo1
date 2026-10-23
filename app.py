import os, secrets, sqlite3, hashlib, hmac, json, re, unicodedata
from datetime import datetime, timezone
from pathlib import Path
from functools import wraps
from flask import Flask, request, redirect, session, render_template_string, send_from_directory, flash, abort
from werkzeug.utils import secure_filename
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
UPLOAD_DIR = ROOT / "private_uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
DB_PATH = ROOT / "app.db"
MAP_PATH = ROOT / "mapa_consulta_base.html"
BACKUP_MAP_PATH = ROOT / "mapa_consulta_base.initial.html"
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}

app = Flask(__name__)
app.secret_key = os.environ.get("APP_SECRET_KEY")
if not app.secret_key:
    raise RuntimeError("Configura APP_SECRET_KEY como variable de entorno antes de iniciar.")

VIEWER_PASSWORD = os.environ.get("VIEWER_PASSWORD")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD")
if not VIEWER_PASSWORD or not ADMIN_PASSWORD:
    raise RuntimeError("Configura VIEWER_PASSWORD y ADMIN_PASSWORD como variables de entorno.")
if VIEWER_PASSWORD == ADMIN_PASSWORD:
    raise RuntimeError("La contraseña de consulta y la administrativa deben ser diferentes.")

app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "true").lower() == "true",
    MAX_CONTENT_LENGTH=20 * 1024 * 1024,
)

def init_db():
    with sqlite3.connect(DB_PATH) as con:
        con.execute("""CREATE TABLE IF NOT EXISTS uploads (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            original_name TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            uploaded_at TEXT NOT NULL,
            uploaded_by TEXT NOT NULL,
            status TEXT NOT NULL
        )""")
init_db()

def role_required(*roles):
    def deco(fn):
        @wraps(fn)
        def wrapped(*args, **kwargs):
            if session.get("role") not in roles:
                return redirect("/login")
            return fn(*args, **kwargs)
        return wrapped
    return deco

LOGIN = """<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Consulta de Vivienda Rural</title><style>
body{font-family:Arial,sans-serif;background:#eef3f5;margin:0;display:grid;min-height:100vh;place-items:center}
.card{background:white;padding:2rem;border-radius:16px;box-shadow:0 8px 30px #1232;width:min(90%,390px)}
h1{font-size:1.35rem;color:#174a3b}label{display:block;margin:1rem 0 .4rem}input,button{box-sizing:border-box;width:100%;padding:.8rem;border-radius:8px;border:1px solid #aaa;font-size:1rem}button{margin-top:1rem;background:#176b52;color:white;border:0;cursor:pointer}
small{color:#555}.msg{color:#a11}</style><div class="card"><h1>Vivienda Rural</h1><p>Demostración con datos ficticios</p>{% for m in get_flashed_messages() %}<p class="msg">{{m}}</p>{% endfor %}
<form method="post"><label for="password">Contraseña</label><input id="password" name="password" type="password" required autocomplete="current-password"><button>Ingresar</button></form><p><small>El acceso administrativo requiere una credencial distinta.</small></p></div></html>"""

ADMIN = """<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Administración</title>
<style>body{font-family:Arial;margin:2rem auto;padding:0 1rem;max-width:850px;color:#19352d}section{border:1px solid #ccd8d3;padding:1.2rem;border-radius:12px;margin:1rem 0}input,button{padding:.7rem;margin:.3rem 0}button{background:#176b52;color:white;border:0;border-radius:6px}li{margin:.5rem 0}.warn{background:#fff5d8;padding:1rem}</style>
<h1>Administración · Vivienda Rural</h1><p>Sesión administrativa activa. <a href="/logout">Cerrar sesión</a></p>
<div class="warn"><b>Importante:</b> Esta es una demostración con datos ficticios. La carga de Excel intenta actualizar el mapa a partir de la hoja “Resumen por Municipio”. No cargues información real ni confidencial en esta demo.</div>
<section><h2>Cargar Excel</h2><form method="post" action="/admin/upload" enctype="multipart/form-data"><input type="file" name="workbook" accept=".xlsx,.xlsm" required><button type="submit">Cargar y registrar archivo</button></form>
<p>Máximo 20 MB. Se admiten archivos .xlsx y .xlsm.</p></section>
<section><h2>Historial de cargas</h2><ul>{% for u in uploads %}<li>{{u[0]}} — {{u[1]}} — {{u[2]}} — {{u[3]}}</li>{% else %}<li>No hay archivos cargados.</li>{% endfor %}</ul></section>
</html>"""

def norm_key(value):
    value = unicodedata.normalize("NFD", str(value or ""))
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^A-Z0-9]+", " ", value.upper()).strip()

def as_num(value):
    if value is None or value == "": return 0
    try: return float(value) if isinstance(value, float) else int(value)
    except (ValueError, TypeError):
        try: return float(str(value).replace(".", "").replace(",", "."))
        except (ValueError, TypeError): return 0

def workbook_to_records(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    if "Resumen por Municipio" not in wb.sheetnames:
        raise ValueError("Falta la hoja ‘Resumen por Municipio’.")
    ws = wb["Resumen por Municipio"]
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True), 1):
        vals = [norm_key(v) for v in row]
        if "ESQUEMA" in vals and "DEPARTAMENTO" in vals and "MUNICIPIO" in vals:
            header_row = i; headers = vals; break
    if not header_row: raise ValueError("No se identificaron encabezados de Esquema, Departamento y Municipio.")
    def col(*names):
        for name in names:
            n = norm_key(name)
            if n in headers: return headers.index(n)
        return None
    ix = {
      "scheme":col("Esquema"), "dep":col("Departamento"), "mun":col("Municipio"),
      "base":col("Cupos base (Fonvivienda / ajustados)","Cupos base"),
      "assigned":col("Asignados (Base de Datos, único)","Asignados"),
      "enabled":col("Habilitados"), "process":col("En proceso de registro"),
      "registered":col("Hogar registrado"), "not_enabled":col("No habilitado"),
      "diagnostic":col("No habilitado por diagnóstico"), "loss":col("Pérdida de ejecutoriedad"),
      "renounce":col("Renuncia al subsidio"), "total":col("Total cupos reportados"),
      "gap":col("Brecha (cupo base - asignados)","Brecha"),
      "vr":col("Cupos contrapartida VR"), "sgr":col("Cupos contrapartida SGR"),
      "alert":col("Observaciones / alertas")
    }
    if any(ix[k] is None for k in ("scheme","dep","mun","base","assigned","enabled")):
        raise ValueError("Faltan columnas obligatorias en Resumen por Municipio.")
    def get(row,key): return row[ix[key]] if ix[key] is not None and ix[key] < len(row) else None
    def clean(v): return str(v).strip() if v is not None else ""
    def key(v): return norm_key(v).replace(" ", " ")
    out=[]
    for row in ws.iter_rows(min_row=header_row+1, values_only=True):
        scheme,dep,mun=clean(get(row,"scheme")),clean(get(row,"dep")),clean(get(row,"mun"))
        if not scheme or not dep or not mun: continue
        alert=clean(get(row,"alert"))
        if alert.lower() in ("nan","none"): alert=""
        base=as_num(get(row,"base")); assigned=as_num(get(row,"assigned"))
        out.append({"esquema":scheme,"departamento":dep,"depKey":key(dep),"municipio":mun,"munKey":key(mun),
          "base":base,"asignados":assigned,"habilitados":as_num(get(row,"enabled")),"proceso":as_num(get(row,"process")),
          "registrado":as_num(get(row,"registered")),"no_habilitado":as_num(get(row,"not_enabled")),
          "diagnostico":as_num(get(row,"diagnostic")),"perdida":as_num(get(row,"loss")),"renuncia":as_num(get(row,"renounce")),
          "total_reportado":as_num(get(row,"total")),"brecha":as_num(get(row,"gap")) if get(row,"gap") is not None else base-assigned,
          "contrapartida_vr":as_num(get(row,"vr")),"contrapartida_sgr":as_num(get(row,"sgr")),"alerta":alert})
    cutoff="no identificado"
    for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,5),values_only=True):
        for val in row:
            txt=clean(val)
            m=re.search(r"Corte\s+(.+?)(?:\s*\||$)",txt,re.I)
            if m: cutoff=m.group(1).strip(); break
        if cutoff!="no identificado": break
    wb.close()
    return out,cutoff

@app.get("/login")
def login_get():
    return render_template_string(LOGIN)

@app.post("/login")
def login_post():
    supplied = request.form.get("password", "")
    if hmac.compare_digest(supplied, VIEWER_PASSWORD):
        session.clear(); session["role"] = "viewer"
        return redirect("/")
    if hmac.compare_digest(supplied, ADMIN_PASSWORD):
        session.clear(); session["role"] = "admin"
        return redirect("/admin")
    flash("Contraseña incorrecta.")
    return redirect("/login")

@app.get("/logout")
def logout():
    session.clear()
    return redirect("/login")

@app.get("/")
@role_required("viewer", "admin")
def viewer():
    # The map is served only after successful server-side authentication.
    return send_from_directory(ROOT, "mapa_consulta_base.html")

@app.get("/admin")
@role_required("admin")
def admin():
    with sqlite3.connect(DB_PATH) as con:
        uploads = con.execute("SELECT original_name, uploaded_at, uploaded_by, status FROM uploads ORDER BY id DESC LIMIT 20").fetchall()
    return render_template_string(ADMIN, uploads=uploads)

@app.post("/admin/upload")
@role_required("admin")
def upload():
    f = request.files.get("workbook")
    if not f or not f.filename:
        flash("Selecciona un archivo Excel.")
        return redirect("/admin")
    original = secure_filename(f.filename)
    ext = Path(original).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        abort(400, "Formato no permitido. Usa .xlsx o .xlsm.")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    stored = f"{stamp}_{secrets.token_hex(6)}_{original}"
    f.save(UPLOAD_DIR / stored)
    uploaded_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    try:
        records, cutoff = workbook_to_records(UPLOAD_DIR / stored)
        if not records:
            raise ValueError("La hoja Resumen por Municipio no contiene filas de datos válidas.")
        html = MAP_PATH.read_text(encoding="utf-8")
        pattern = re.compile(r"let records = \[.*?\];", re.S)
        replacement = "let records = " + json.dumps(records, ensure_ascii=False, separators=(",", ":")) + ";"
        updated, count = pattern.subn(lambda _m: replacement, html, count=1)
        if count != 1:
            raise ValueError("No se encontró el bloque de datos del mapa; no se publicaron cambios.")
        updated = re.sub(r'Fuente inicial: archivo .*? Esta visualización es de seguimiento y no reemplaza la validación oficial de cifras\.',
                         f'Fuente inicial: archivo “{original}”, hoja “Resumen por Municipio”, corte {cutoff}. Esta versión se configura para MEJORAMIENTOS; no incluye vivienda nueva. Los totales se calculan sumando las filas visibles del resumen; al cambiar de esquema o departamento, los indicadores se recalculan. Cartografía: MGN 2025, DANE; simplificación y distribución de mapa: ytolosa/mapas-colombia, licencia CC BY 4.0. Esta visualización es de seguimiento y no reemplaza la validación oficial de cifras.', updated)
        # Save only after workbook validation and successful transformation.
        MAP_PATH.write_text(updated, encoding="utf-8")
        status = f"Publicado: {len(records)} filas; corte {cutoff}"
    except Exception as exc:
        status = "Error de validación: " + str(exc)[:240]
    with sqlite3.connect(DB_PATH) as con:
        con.execute("INSERT INTO uploads(original_name,stored_name,uploaded_at,uploaded_by,status) VALUES(?,?,?,?,?)",
                    (original, stored, uploaded_at, "admin", status))
    return redirect("/admin")

@app.get("/health")
def health():
    return {"status": "ok"}

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8000")), debug=False)
